import docx
import re
from openpyxl import Workbook

def parse_word_structure(file_path):
    doc = docx.Document(file_path)

    structure = {}

    current_category = ""
    current_subcategories = []
    current_content = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()

        # 匹配标题编号
        match = re.match(r'^(\d+(\.\d+)*)\s+', text)
        if match:
            heading_number = match.group(1)

            # 处理之前的数据
            if current_category:
                structure[current_category] = {'Subcategories': current_subcategories, 'Content': current_content}

            # 更新当前分类
            current_category = heading_number
            current_subcategories = [text]
            current_content = []

        elif current_category and text:
            # 区分子分类和正文内容
            if text.startswith(current_category):
                current_subcategories.append(text)
            else:
                current_content.append(text)

    # 处理最后一组数据
    if current_category:
        structure[current_category] = {'Subcategories': current_subcategories, 'Content': current_content}

    return structure

def print_structure(structure,excel):
    wb = Workbook()
    ws = wb.worksheets[0]
    ws.append(['问题','回答'])
    subcategories = ''
    for category, data in structure.items():
        cc = str(data['Subcategories']).split(maxsplit=1)[0]
        c = cc.count('.')
        if c>0 and not cc.isdigit():
            subcategories = "".join(str(data['Subcategories']).replace(cc,'').replace('''']''',''))
            #print(str(data['Subcategories']).replace(cc,'').replace('''']''',''))
            content = "   ".join(data['Content'])
        
            ws.append([subcategories,content])
    wb.save(excel)

# 示例使用
fn_word = 'test.docx'  # 请替换为实际的 Word 文档路径
fn_excel = fn_word[:-4]+'.xlsx'

document_structure = parse_word_structure(fn_word)
print_structure(document_structure,fn_excel)

